"""Desktop-selected files, durable snapshots and automatic evidence preparation."""
import io,json,stat
from pathlib import Path
from memorive_folder_management.policy import windows_io_path
from .store import digest,uid,now,packed

METHODS=frozenset({'memo.attach_files','memo.attachment_status','memo.thread_configure','memo.material_preview','memo.material_remove'})
SUPPORTED={'.pdf','.md','.txt','.json','.docx','.png','.jpg','.jpeg','.webp'}
IMAGES={'.png','.jpg','.jpeg','.webp'}

class Attachments:
    def __init__(self,workspace):
        self.w=workspace;self.store=workspace.store
        self.root=self.store.root/'attachments';windows_io_path(self.root).mkdir(exist_ok=True)
        import threading
        self._locks={};self._locks_guard=threading.Lock()

    def profile(self,thread):
        return thread.get('profile_ref',self.w.settings()['profile_ref'])

    def configure(self,thread_id,expected_revision,changes):
        if not isinstance(changes,dict) or not changes or set(changes)-{'profile_ref','artifact_ids','scope','include_project'}:
            raise ValueError('THREAD_CONFIG_INVALID')
        with self.store.tx() as db:
            thread=self.w.conversations._get(thread_id,expected_revision,db)
            if thread.get('archived'):raise ValueError('THREAD_ARCHIVED')
            if 'profile_ref' in changes:
                ref=changes['profile_ref']
                if ref is not None and not any(o['profile_ref']==ref and o['eligible'] for o in self.w.catalog()):
                    raise ValueError('CHAT_MODEL_UNAVAILABLE')
            if 'artifact_ids' in changes:
                ids=changes['artifact_ids']
                if not isinstance(ids,list) or len(ids)>32 or any(not isinstance(i,str) for i in ids):raise ValueError('MATERIAL_SELECTION_INVALID')
                for i in ids:
                    item=self.store.get('artifact',i,db=db)
                    if not item or item['project']!=thread['project']:raise ValueError('MATERIAL_OUTSIDE_PROJECT')
                removed=set(thread.get('artifact_ids',[]))-set(ids)
                excluded=(set(thread.get('excluded_artifact_ids',[]))|removed)-set(ids)
                changes=dict(changes,artifact_ids=list(dict.fromkeys(ids)),excluded_artifact_ids=sorted(excluded))
            if 'include_project' in changes:
                if type(changes['include_project']) is not bool:raise ValueError('MATERIAL_SCOPE_INVALID')
                changes=dict(changes,scope='project' if changes['include_project'] else 'conversation')
            if 'scope' in changes:changes=dict(changes,include_project=changes['scope']=='project')
            if changes.get('scope',thread.get('scope','project')) not in {'project','conversation'}:raise ValueError('MATERIAL_SCOPE_INVALID')
            return self.store.put('thread',thread_id,thread['project'],dict(thread,**changes),expected=expected_revision,db=db)

    def attach(self,thread_id,paths):
        thread=self.w.thread_get(thread_id)
        with self.store.tx() as db:self.w.conversations._get(thread_id,thread['revision'],db)
        if thread.get('archived'):raise ValueError('THREAD_ARCHIVED')
        if not isinstance(paths,list) or not 1<=len(paths)<=32:raise ValueError('ATTACHMENT_FILES_REQUIRED')
        job_id=uid('attach_');items=[];ref=self.profile(thread)
        for value in paths:
            try:
                if not isinstance(value,str):raise ValueError('ATTACHMENT_PATH_INVALID')
                path=Path(value);io_path=windows_io_path(path)
                if not path.is_absolute() or not io_path.is_file() or io_path.is_symlink():raise ValueError('ATTACHMENT_PATH_INVALID')
                if getattr(io_path.lstat(),'st_file_attributes',0)&getattr(stat,'FILE_ATTRIBUTE_REPARSE_POINT',0x400):raise ValueError('SOURCE_LINK_REJECTED')
                if path.suffix.lower() not in SUPPORTED:raise ValueError('ATTACHMENT_FORMAT_UNSUPPORTED')
                if io_path.stat().st_size>64*1024*1024:raise ValueError('SOURCE_TOO_LARGE')
                raw=io_path.read_bytes()
                if not raw:raise ValueError('ATTACHMENT_EMPTY')
                if path.suffix.lower()=='.json':
                    candidate=json.loads(raw.decode('utf-8-sig'))
                    schema=candidate.get('schema','') if isinstance(candidate,dict) else ''
                    if isinstance(schema,str) and schema.startswith('memo-source/'):
                        result=self.w.skill.import_source(thread_id,path)
                        items.append(dict({k:result[k] for k in ('id','status','reused','source_sha256','conversion_called','warnings')},name=path.name));continue
                    if isinstance(schema,str) and schema.startswith('memo-research-answer/'):
                        if schema!='memo-research-answer/1':raise ValueError('SKILL_ANSWER_SCHEMA_UNSUPPORTED')
                        binding=self.store.get('handoff',candidate.get('handoff_id',''))
                        if not binding or binding['thread_id']!=thread_id:raise ValueError('HANDOFF_THREAD_MISMATCH')
                        result=self.w.skill.answer(**{k:v for k,v in candidate.items() if k!='schema'})
                        items.append(dict(name=path.name,status='READY',answer_return=result));continue
                content_hash=digest(raw);identity='art_'+digest([thread['project'],'attachment',thread_id if thread['temporary'] else '',content_hash])[:32]
                reused=self.w.skill.reuse_original(thread_id,content_hash)
                if reused:
                    items.append(dict(reused,name=path.name));continue
                dest=self.root/(content_hash+path.suffix.lower())
                io_dest=windows_io_path(dest)
                if io_dest.exists():
                    if digest(io_dest.read_bytes())!=content_hash:raise ValueError('ATTACHMENT_SNAPSHOT_CONFLICT')
                else:
                    with io_dest.open('xb') as f:f.write(raw)
                with self.store.tx() as db:
                    old=self.store.get('artifact',identity,db=db)
                    reusable=bool(old and old['state']=='active')
                    pending=bool(old and old['state']=='preparing')
                    if not reusable and not pending:
                        self.store.put('artifact',identity,thread['project'],{'project':thread['project'],'title':(old or {}).get('title',path.name),
                            'path':str(dest),'root':str(self.root),'content_hash':content_hash,'kind':'attachment','state':'preparing',
                            'document_id':'doc_'+content_hash,'temporary_thread':thread_id if thread['temporary'] else None,'original_name':(old or {}).get('original_name',path.name),'created_at':now(),'chunks':0,'profile_ref':ref},db=db)
                    fresh=self.store.get('thread',thread_id,db=db)
                    if not fresh or fresh.get('archived'):raise ValueError('THREAD_ARCHIVED')
                    ids=list(dict.fromkeys(fresh.get('artifact_ids',[])+[identity]))
                    if len(ids)>32:raise ValueError('ATTACHMENT_SCOPE_LIMIT')
                    self.store.put('thread',thread_id,thread['project'],dict(fresh,artifact_ids=ids,excluded_artifact_ids=[i for i in fresh.get('excluded_artifact_ids',[]) if i!=identity]),db=db)
                items.append({'id':identity,'name':path.name,'status':'READY' if reusable else 'QUEUED','reused':reusable})
            except (ValueError,OSError,TypeError,UnicodeError) as exc:
                items.append({'name':Path(str(value)).name,'status':'ERROR','error':str(exc)[:140]})
        job=self.store.put('attachment_job',job_id,thread['project'],{'thread_id':thread_id,'project':thread['project'],
            'profile_ref':ref,'status':'QUEUED','items':items,'created_at':now()})
        with self.w.lock:
            self.w.futures[job_id]=self.w.pool.submit(self._prepare,job_id)
        return job

    def status(self,job_id):
        job=self.store.get('attachment_job',job_id)
        if not job:raise ValueError('ATTACHMENT_JOB_NOT_FOUND')
        return job

    def _prepare(self,job_id):
        job=self.status(job_id);job['status']='RUNNING';self.store.put('attachment_job',job_id,job['project'],job)
        for item in job['items']:
            if item['status']!='QUEUED':continue
            with self._locks_guard:
                import threading
                lock=self._locks.setdefault(item['id'],threading.Lock())
            lock.acquire()
            artifact=self.store.get('artifact',item['id'])
            try:
                if artifact['state']=='active':
                    item['status']='READY';continue
                pages,receipts=self._pages(artifact,job['profile_ref'],job_id)
                if not any(text.strip() for _,text in pages):raise ValueError('ATTACHMENT_TEXT_EMPTY')
                chunks=[]
                for page,text in pages:
                    lines=text.splitlines();buf=[];size=0;start=1
                    for end,line in enumerate(lines,1):
                        buf.append(line);size+=len(line)+1
                        if size>=1800 or end==len(lines):
                            content='\n'.join(buf)
                            if content.strip():
                                cid='ev_'+digest([artifact['id'],artifact['content_hash'],page,start,end,digest(content)])[:32]
                                chunks.append((cid,artifact['id'],artifact['project'],content,start,end,page,digest(content)))
                            buf=[];size=0;start=end+1
                extraction=self.root/(artifact['content_hash']+'.'+digest(pages)[:16]+'.pages.json')
                encoded=packed(pages).encode('utf-8')
                io_extraction=windows_io_path(extraction)
                if not io_extraction.exists():
                    with io_extraction.open('xb') as f:f.write(encoded)
                with self.store.tx() as db:
                    self.store.put('artifact',artifact['id'],artifact['project'],dict(artifact,state='active',indexed_at=now(),
                        chunks=len(chunks),extraction_path=str(extraction),extraction_hash=digest(encoded),
                        page_count=len(pages),vision_receipts=receipts,error=None),db=db)
                    db.execute('DELETE FROM chunks WHERE artifact_id=?',(artifact['id'],))
                    db.executemany('INSERT INTO chunks VALUES(?,?,?,?,?,?,?,?)',chunks)
                item['status']='READY'
            except Exception as exc:
                item.update(status='ERROR',error=str(exc)[:180])
                self.store.put('artifact',artifact['id'],artifact['project'],dict(artifact,state='error',error=item['error']))
            finally:lock.release()
            self.store.put('attachment_job',job_id,job['project'],job)
        job['status']='COMPLETE' if all(x['status']=='READY' for x in job['items']) else 'PARTIAL'
        self.store.put('attachment_job',job_id,job['project'],job)

    def _pages(self,artifact,profile_ref,job_id):
        path=Path(artifact['path']);raw=windows_io_path(path).read_bytes();suffix=path.suffix.lower();receipts=[]
        def vision(data,mime,page):
            if not profile_ref:raise ValueError('ATTACHMENT_VISION_MODEL_REQUIRED')
            if not self.w.vision:raise ValueError('ATTACHMENT_VISION_UNAVAILABLE')
            result=self.w.vision(profile_ref=profile_ref,image_bytes=data,mime_type=mime,
                job_id='research-'+digest([job_id,artifact['id'],page])[:24])
            receipt=result.get('execution_receipt') or {}
            if result.get('status')!='PASS' or receipt.get('status')!='PASS':raise ValueError(result.get('reason') or 'ATTACHMENT_VISION_FAILED')
            if not __import__('memorive_settings.cli_templates',fromlist=['receipt_evidence_complete']).receipt_evidence_complete(receipt):raise ValueError('ATTACHMENT_VISION_RECEIPT_REQUIRED')
            receipts.append(receipt)
            return result['text']
        if suffix=='.pdf':
            import fitz
            with fitz.open(stream=raw,filetype='pdf') as doc:
                if doc.needs_pass:raise ValueError('ATTACHMENT_PASSWORD_REQUIRED')
                pages=[]
                for page in doc:
                    text=page.get_text()
                    if not text.strip() or (len(text.strip())<32 and page.get_images()):
                        pix=page.get_pixmap(matrix=fitz.Matrix(1.4,1.4),alpha=False)
                        text=vision(pix.tobytes('png'),'image/png',page.number+1)
                    pages.append((page.number+1,text))
                return pages,receipts
        if suffix in IMAGES:
            from PIL import Image,ImageOps
            with Image.open(io.BytesIO(raw)) as im:
                im=ImageOps.exif_transpose(im).convert('RGB');im.thumbnail((2400,2400));output=io.BytesIO();im.save(output,format='PNG')
            return [(1,vision(output.getvalue(),'image/png',1))],receipts
        if suffix=='.docx':
            from docx import Document
            doc=Document(io.BytesIO(raw))
            return [(None,'\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows]))],receipts
        text=raw.decode('utf-8-sig')
        if suffix=='.json':text=json.dumps(json.loads(text),ensure_ascii=False,indent=2)
        return [(None,text)],receipts

    def preview(self,project,artifact_id):
        a=self.store.get('artifact',artifact_id)
        if not a or a['project']!=project:raise ValueError('MATERIAL_OUTSIDE_PROJECT')
        with self.store.tx() as db:
            refs=[self.w.index.read(r[0],project,db=db) for r in db.execute('SELECT id FROM chunks WHERE artifact_id=? ORDER BY page,line_start LIMIT 6',(artifact_id,))]
        return {'artifact':a,'evidence':refs}

    def remove(self,thread_id,expected_revision,artifact_ids):
        thread=self.w.thread_get(thread_id)
        if not isinstance(artifact_ids,list) or not artifact_ids or any(not isinstance(i,str) for i in artifact_ids):raise ValueError('MATERIAL_SELECTION_INVALID')
        with self.store.tx() as db:
            thread=self.w.conversations._get(thread_id,expected_revision,db)
            if thread.get('archived'):raise ValueError('THREAD_ARCHIVED')
            for identity in artifact_ids:
                a=self.store.get('artifact',identity,db=db)
                if not a or a['project']!=thread['project']:raise ValueError('MATERIAL_OUTSIDE_PROJECT')
            # Removing from a conversation never destroys the original or another conversation's evidence.
            value=dict(thread,artifact_ids=[i for i in thread.get('artifact_ids',[]) if i not in artifact_ids],
                       excluded_artifact_ids=sorted(set(thread.get('excluded_artifact_ids',[]))|set(artifact_ids)))
            self.store.event('MATERIALS_REMOVED_FROM_THREAD',thread_id,{'artifact_ids':artifact_ids},db)
            return self.store.put('thread',thread_id,thread['project'],value,expected=expected_revision,db=db)

    def scope_ids(self,thread,*,db=None):
        selected=list(thread.get('artifact_ids',[]))
        if thread.get('include_project',thread.get('scope','project')=='project'):
            selected+= [a['id'] for a in self.store.list('artifact',thread['project'],db=db) if a['state']=='active' and not a.get('temporary_thread')]
        return list(dict.fromkeys(i for i in selected if i not in thread.get('excluded_artifact_ids',[])))

    def call(self,method,params):
        return {'memo.attach_files':self.attach,'memo.attachment_status':self.status,
                'memo.thread_configure':self.configure,'memo.material_preview':self.preview,'memo.material_remove':self.remove}[method](**params)

